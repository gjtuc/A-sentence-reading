"""design/295 — VML imagedata figures and caption-stub sentence merge."""

from __future__ import annotations

import subprocess
import sys
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from PIL import Image

from raster_floor import png_over_docx_min
from sentence_reading.docx.extract import _MIN_BYTES, extract_figures, figure_source_census
from sentence_reading.pdf.sentences import split_into_sentences

ROOT = Path(__file__).resolve().parents[1]


def _replace_blip_with_vml(paragraph) -> None:
    blip = next(el for el in paragraph._element.iter() if str(el.tag).endswith("}blip"))
    rid = blip.get(qn("r:embed"))
    drawing = next(el for el in paragraph._element.iter() if str(el.tag).endswith("}drawing"))
    drawing.getparent().remove(drawing)
    pict = parse_xml(
        "<w:r xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\" "
        "xmlns:v=\"urn:schemas-microsoft-com:vml\" "
        "xmlns:r=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships\">"
        f"<w:pict><v:shape><v:imagedata r:id=\"{rid}\"/></v:shape></w:pict></w:r>"
    )
    paragraph._element.append(pict)


def test_vml_imagedata_pairs_with_next_caption(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run().add_picture(BytesIO(png_over_docx_min()))
    _replace_blip_with_vml(p)
    doc.add_paragraph("Fig. S1. Rietveld refinement of the sample.")
    path = tmp_path / "si.docx"
    doc.save(path)

    figs = extract_figures(path, doc_role="supplementary")
    assert len(figs) == 1
    assert figs[0].caption.startswith("Fig. S1")
    assert figs[0].slot_key == "fig:s1"
    census = figure_source_census(path)
    assert census["blip_n"] == 0
    assert census["imagedata_n"] == 1
    assert census["vml_unseen_n"] == 0


def test_main_docx_does_not_stamp_si_slot(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run().add_picture(BytesIO(png_over_docx_min()))
    _replace_blip_with_vml(p)
    doc.add_paragraph("Fig. S2. Not a main slot.")
    path = tmp_path / "main.docx"
    doc.save(path)
    figs = extract_figures(path, doc_role="main")
    assert len(figs) == 1
    assert figs[0].slot_key == ""


def test_drawingml_blip_still_extracts(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph().add_run().add_picture(BytesIO(png_over_docx_min()))
    doc.add_paragraph("Fig. 1. DrawingML still pairs.")
    path = tmp_path / "main.docx"
    doc.save(path)
    figs = extract_figures(path)
    assert len(figs) == 1
    assert "Fig. 1" in figs[0].caption


def test_caption_label_stays_with_body() -> None:
    text = (
        "Ca 2p\n\nFig. S3.\n\nNormalized ECR response curve at 700 C.\n\n"
        "Fig. S4.\n\nTafel plots of the electrode at 700 C."
    )
    cards = [s.text for s in split_into_sentences(text)]
    assert cards == [
        "Ca 2p",
        "Fig. S3. Normalized ECR response curve at 700 C.",
        "Fig. S4. Tafel plots of the electrode at 700 C.",
    ]


def test_prose_fig_period_still_splits() -> None:
    text = "as shown in Fig. S1. The rate increased."
    cards = [s.text for s in split_into_sentences(text)]
    assert cards == ["as shown in Fig. S1.", "The rate increased."]


def test_initial_does_not_split() -> None:
    cards = [s.text for s in split_into_sentences("Liping R. China")]
    assert cards == ["Liping R. China"]
    two = [s.text for s in split_into_sentences("The end. Next claim starts.")]
    assert two == ["The end.", "Next claim starts."]


def test_png_fixture_clears_extract_floor() -> None:
    blob = png_over_docx_min()
    assert len(blob) >= _MIN_BYTES
    # The miss: a solid 40x40 compresses under the floor. Do not guess pixels.
    tiny = BytesIO()
    Image.new("RGB", (40, 40), (180, 20, 20)).save(tiny, format="PNG")
    assert len(tiny.getvalue()) < _MIN_BYTES


def test_vml_tests_do_not_guess_pixel_size() -> None:
    text = Path(__file__).read_text(encoding="utf-8")
    assert "png_over_docx_min" in text
    assert "Image.new" not in text.split("def test_png_fixture_clears_extract_floor")[0]


def test_replay_script_gates_fig_and_stub(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run().add_picture(BytesIO(png_over_docx_min()))
    _replace_blip_with_vml(p)
    doc.add_paragraph("Fig. S4.")
    doc.add_paragraph("Tafel plots of the electrode at 700 C.")
    path = tmp_path / "si.docx"
    doc.save(path)
    proc = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "replay_docx_extract.py"),
            str(path),
            "--min-fig",
            "1",
            "--max-stub",
            "0",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert proc.returncode == 0, proc.stdout + proc.stderr


def test_design_295_indexed() -> None:
    readme = (ROOT / "docs/design/README.md").read_text(encoding="utf-8")
    assert "295 |" in readme
    assert (ROOT / "docs/design/295-docx-vml-caption-split.md").is_file()
