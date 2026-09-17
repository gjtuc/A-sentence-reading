"""Title verify for the import switch. Not used on folder open.

The phone fills the wide-box title locally. This path runs only when the
user flips the switch and the other side is still empty. Repeat 1000-character
windows until a title verifies, then stop. Return the extracted sentence, not
the Crossref display title.
"""

from __future__ import annotations

import json
import re
import tempfile
import urllib.parse
import urllib.request
from pathlib import Path

from sentence_reading.llm.crossref_resolve import _mailto

_SYSTEM = """Extract the article title and author names from an academic paper excerpt.
Return JSON only: {"title":"", "authors":["Family, Given"], "stop":""}
title is the article title, not the journal, not a running header, not a section heading.
authors are people named as authors, not affiliations.
stop is "section" if the excerpt is mainly Results, Conclusion, or References with no article title.
stop is "none" if there is no article title on this excerpt.
Otherwise stop is "".
"""

_MAX_PAGES = 8
_MAX_STEPS = 8


def text_windows(path: Path) -> list[tuple[int, str]]:
    pages = _pages(path)
    windows: list[tuple[int, str]] = []
    for page_i, blob in enumerate(pages, 1):
        blob = blob.replace("\x00", " ")
        start = 0
        step = 0
        while start < len(blob) and step < _MAX_STEPS:
            chunk = blob[start : start + 1000].strip()
            if len(chunk) >= 40:
                windows.append((page_i, chunk))
            start += 1000
            step += 1
            if start >= len(blob):
                break
    return windows


def _pages(path: Path) -> list[str]:
    if path.suffix.lower() == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts: list[str] = []
        for para in doc.paragraphs:
            t = " ".join((para.text or "").split())
            if t:
                parts.append(t)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    t = " ".join((cell.text or "").split())
                    if t and t not in parts:
                        parts.append(t)
        return ["\n".join(parts)]
    import pymupdf

    doc = pymupdf.open(path)
    try:
        return [(doc[i].get_text("text") or "") for i in range(min(_MAX_PAGES, doc.page_count))]
    finally:
        doc.close()


def parse_extract(raw: str) -> dict:
    text = (raw or "").strip()
    data: dict | None = None
    try:
        loaded = json.loads(text)
        if isinstance(loaded, dict):
            data = loaded
    except json.JSONDecodeError:
        data = _salvage(text)
    if not data:
        return {"title": "", "authors": [], "stop": "parse"}
    authors = data.get("authors") or []
    if isinstance(authors, str):
        authors = [authors]
    return {
        "title": str(data.get("title") or "").strip()[:240],
        "authors": [str(a).strip() for a in authors if str(a).strip()][:12],
        "stop": str(data.get("stop") or "").strip().lower(),
    }


def _salvage(raw: str) -> dict | None:
    title_m = re.search(r'"title"\s*:\s*"((?:\\.|[^"\\])*)"', raw)
    if not title_m:
        return None
    title = json.loads(f'"{title_m.group(1)}"')
    authors = re.findall(r'"([A-Za-z][^"]{1,80})"', raw[title_m.end() :])
    return {"title": title, "authors": authors[:8], "stop": ""}


def tokens(raw: str) -> set[str]:
    return {t for t in re.split(r"[^A-Za-z]+", raw.lower()) if len(t) >= 3}


def _word(token: str, blob: str) -> bool:
    if len(token) < 3:
        return False
    return re.search(rf"(?<![a-z]){re.escape(token)}(?![a-z])", blob) is not None


def overlap_file(hit_authors: list[dict], blob: str) -> list[str]:
    low = blob.lower()
    matched: list[str] = []
    for person in hit_authors:
        family = str(person.get("family") or "").lower()
        if not _word(family, low):
            continue
        given_toks = tokens(str(person.get("given") or ""))
        if given_toks and not any(_word(t, low) for t in given_toks):
            continue
        matched.append(family)
    return sorted(set(matched))


def crossref_hits(title: str, author: str) -> list[dict]:
    q = urllib.parse.urlencode(
        {
            "query.bibliographic": title[:180],
            "query.author": author[:80],
            "rows": "5",
            "mailto": _mailto(),
        }
    )
    req = urllib.request.Request(
        "https://api.crossref.org/works?" + q,
        headers={
            "User-Agent": f"A-sentence-reading title-verify (mailto:{_mailto()})",
            "Accept": "application/json",
        },
    )
    with urllib.request.urlopen(req, timeout=20) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    out = []
    for item in data.get("message", {}).get("items") or []:
        names = []
        for person in item.get("author") or []:
            names.append(
                {
                    "given": str(person.get("given") or ""),
                    "family": str(person.get("family") or ""),
                }
            )
        out.append(
            {
                "score": float(item.get("score") or 0),
                "authors": names[:12],
                "doi": str(item.get("DOI") or ""),
            }
        )
    return out


def pick_hit(title: str, authors: list[str], blob: str, search=crossref_hits) -> dict:
    author_q = ""
    if authors:
        parts = re.split(r"[, ]+", authors[0])
        author_q = parts[0] if parts else authors[0]
    best = None
    for hit in search(title, author_q):
        names = overlap_file(hit["authors"], blob)
        if not names:
            continue
        if best is None or hit["score"] > best["score"]:
            best = {**hit, "overlap": names}
    return best or {}


def gemini_extract(chunk: str) -> dict:
    from google import genai
    from google.genai import types
    from sentence_reading.llm.env import gemini_api_key, gemini_model

    client = genai.Client(api_key=gemini_api_key())
    response = client.models.generate_content(
        model=gemini_model(),
        contents=chunk,
        config=types.GenerateContentConfig(
            system_instruction=_SYSTEM,
            temperature=0.0,
            max_output_tokens=1024,
            response_mime_type="application/json",
        ),
    )
    return parse_extract(getattr(response, "text", None) or "")


def verify_windows(windows: list[tuple[int, str]], blob: str, *, extract=gemini_extract, search=crossref_hits) -> dict:
    tried = 0
    for _page, chunk in windows:
        tried += 1
        extracted = extract(chunk)
        title = str(extracted.get("title") or "").strip()
        stop = str(extracted.get("stop") or "")
        if stop in {"section", "none", "parse"} or not title:
            continue
        hit = pick_hit(title, list(extracted.get("authors") or []), blob, search=search)
        if not hit:
            continue
        return {
            "ok": True,
            "title": title,
            "doi": hit.get("doi") or "",
            "windows": tried,
        }
    return {"ok": False, "title": "", "doi": "", "windows": tried, "error": "no_title"}


def verify_path(path: Path) -> dict:
    windows = text_windows(path)
    blob = "\n".join(chunk for _page, chunk in windows)
    if not windows:
        return {"ok": False, "title": "", "doi": "", "windows": 0, "error": "no_text"}
    try:
        return verify_windows(windows, blob)
    except Exception as exc:  # noqa: BLE001
        return {
            "ok": False,
            "title": "",
            "doi": "",
            "windows": 0,
            "error": type(exc).__name__,
        }


def verify_bytes(raw: bytes, filename: str) -> dict:
    suffix = Path(filename or "paper.pdf").suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        suffix = ".pdf"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(raw)
        path = Path(tmp.name)
    try:
        return verify_path(path)
    finally:
        path.unlink(missing_ok=True)
